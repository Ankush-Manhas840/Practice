"""
Regenerates questions-log.xlsx from progress.json.

Run this any time progress.json changes:
    python scripts/generate_questions_log.py

Then recalc formulas (LibreOffice headless):
    python <xlsx-skill>/scripts/recalc.py questions-log.xlsx
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PROGRESS_JSON = ROOT / "progress.json"
OUTPUT_XLSX = ROOT / "questions-log.xlsx"

DIFFICULTY_FILL = {
    "Easy": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "Medium": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "Hard": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
BODY_FONT = Font(name="Arial")

COLUMNS = [
    ("Title", "title", 42),
    ("Topic", "topic", 22),
    ("Difficulty", "difficulty", 12),
    ("First Attempted", "first_attempted", 16),
    ("Status", "status", 10),
    ("Times Reviewed", "times_reviewed", 14),
    ("Last Reviewed", "last_reviewed", 14),
    ("Next Review", "next_review", 14),
    ("Solution File", "solution_file", 55),
]


def load_solved_questions():
    data = json.loads(PROGRESS_JSON.read_text())
    solved = [
        q for q in data["questions"]
        if q.get("status") == "solved" and q.get("first_attempted")
    ]
    solved.sort(key=lambda q: (q["first_attempted"], q["topic"]))
    return solved, data


def build_questions_sheet(wb, solved):
    ws = wb.active
    ws.title = "Questions"

    for col_idx, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    for row_idx, q in enumerate(solved, start=2):
        for col_idx, (_, key, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=q.get(key))
            cell.font = BODY_FONT
            if key == "difficulty" and q.get(key) in DIFFICULTY_FILL:
                cell.fill = DIFFICULTY_FILL[q[key]]

    return ws


def build_summary_sheet(wb, solved, all_topics):
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14

    title = ws.cell(row=1, column=1, value="DSA Practice Summary")
    title.font = Font(name="Arial", bold=True, size=14)

    ws.cell(row=3, column=1, value="Total Solved").font = Font(name="Arial", bold=True)
    total_cell = ws.cell(row=3, column=2, value="=COUNTA(Questions!A2:A10000)")
    total_cell.font = BODY_FONT

    header_row = 5
    ws.cell(row=header_row, column=1, value="Topic").font = Font(name="Arial", bold=True)
    ws.cell(row=header_row, column=2, value="Count").font = Font(name="Arial", bold=True)
    ws.cell(row=header_row, column=1).fill = HEADER_FILL
    ws.cell(row=header_row, column=2).fill = HEADER_FILL
    ws.cell(row=header_row, column=1).font = HEADER_FONT
    ws.cell(row=header_row, column=2).font = HEADER_FONT

    topic_counts = {}
    for q in solved:
        topic_counts[q["topic"]] = topic_counts.get(q["topic"], 0) + 1
    ordered_topics = sorted(topic_counts, key=lambda t: -topic_counts[t])

    for i, topic in enumerate(ordered_topics, start=1):
        r = header_row + i
        ws.cell(row=r, column=1, value=topic).font = BODY_FONT
        ws.cell(row=r, column=2, value=f'=COUNTIF(Questions!B:B,"{topic}")').font = BODY_FONT

    return ws


def main():
    solved, data = load_solved_questions()
    wb = Workbook()
    build_questions_sheet(wb, solved)
    build_summary_sheet(wb, solved, data.get("topics", []))
    wb.save(OUTPUT_XLSX)
    print(f"Wrote {len(solved)} solved questions to {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
